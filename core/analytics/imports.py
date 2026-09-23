"""CSV/XLSX in memory, explicit nulls and reproducible, isolated row identities."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import unicodedata
import zipfile
from collections import Counter
from importlib.resources import files

from openpyxl import load_workbook

from core.analytics.profiles import evidence_confidence, player_profile
from core.pipeline.metrics import DEFINITIONS, metric_catalog
from core.schemas.imports import ClubImport, IMPORT_FIELDS, POSITIONS, PREVIEW_CELL_CHARACTERS

MAX_FILE_BYTES = 3 * 1024 * 1024
MAX_JSON_BYTES = 8 * 1024 * 1024
# Include transitive bounds so a missing/rejected parent cannot hide an
# impossible subtotal. Shared with the frontend's response validator.
METRIC_BOUNDS = json.loads(
    files("core").joinpath("reference/import_metric_bounds.json").read_text(encoding="utf-8")
)
ALIASES = {
    "player_name": ["player", "name", "nombre", "jugador", "joueur", "nom"],
    "player_id": ["id"],
    "team_name": ["team", "club", "equipo", "equipe", "team_in_tournament"],
    "position_group": ["position", "posicion", "poste", "pos"],
    "minutes": ["min", "mins", "minutos", "minutes_including_added_time"],
    "age": ["edad", "age"],
    "goals": ["goles", "buts"],
    "shots": ["tiros", "tirs"],
    "assists": ["asistencias", "passes_decisives"],
    "passes": ["pases", "passes_attempted"],
    "completed_passes": ["pases_completados", "passes_completed"],
}


def normalized(value: str) -> str:
    text = "".join(c for c in unicodedata.normalize("NFKD", value) if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "_", text.casefold()).strip("_")


def preview_table(content: bytes, filename: str) -> dict:
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("Use a non-empty CSV or XLSX file up to 3 MiB")
    warnings = []
    table = []
    row_numbers = []
    if filename.casefold().endswith(".csv"):
        encoding = "utf-16" if content[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8-sig"
        try:
            decoded = content.decode(encoding)
            dialect = csv.Sniffer().sniff(decoded[:10000], delimiters=",;\t")
        except UnicodeError as exc:
            raise ValueError("CSV must use UTF-8 or UTF-16; export it again with Unicode encoding") from exc
        except csv.Error:
            dialect = csv.excel
        source_rows = csv.reader(io.StringIO(decoded), dialect=dialect, strict=True)
        book = None
    elif filename.casefold().endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if any(i.flag_bits & 1 for i in archive.infolist()):
                    raise ValueError("Password-protected workbooks are not supported")
                if (
                    len(archive.infolist()) > 100
                    or sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024
                ):
                    raise ValueError("Workbook expands beyond the 20 MiB / 100-entry limit")
                if any("vbaProject" in name for name in archive.namelist()):
                    raise ValueError("Macro-enabled workbooks are not supported")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            sheet = book.active
            if sheet is None or not hasattr(sheet, "iter_rows"):
                book.close()
                raise ValueError("Workbook has no readable worksheet")
            if len(book.sheetnames) > 1:
                warnings.append(f"Only active worksheet '{sheet.title}' was read; other sheets were ignored")
            # Bounds are independent of workbook-declared dimensions, which can be forged.
            sheet.reset_dimensions()
            source_rows = sheet.iter_rows(values_only=True)
        except (zipfile.BadZipFile, KeyError, OSError) as exc:
            raise ValueError("Invalid XLSX workbook") from exc
    else:
        raise ValueError("Supported formats are .csv and .xlsx; convert legacy .xls files first")
    formula_count = 0
    try:
        for physical_row, row in enumerate(source_rows, 1):
            if physical_row > 2001:
                raise ValueError("Maximum 2,000 player rows including blank rows")
            cells = [None if v is None else str(v).strip() for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if not any(cells):
                continue
            if len(cells) > 60 or any(c is not None and len(c) > PREVIEW_CELL_CHARACTERS for c in cells):
                raise ValueError(f"Maximum 60 columns and {PREVIEW_CELL_CHARACTERS} characters per cell")
            for i, cell in enumerate(cells):
                if cell and cell.startswith(("=", "+", "@")):
                    cells[i] = None
                    formula_count += 1
            table.append(cells)
            row_numbers.append(physical_row)
            if len(table) > 2001:
                raise ValueError("Maximum 2,000 player rows; split the file by squad or season")
    except csv.Error as exc:
        raise ValueError("Malformed CSV quoting or oversized field") from exc
    except (IndexError, KeyError, OSError) as exc:
        # Read-only XLSX parsing is lazy: broken shared-string/cell references
        # can surface here, after load_workbook has already returned.
        raise ValueError("Invalid XLSX cell or shared-string reference") from exc
    finally:
        if book is not None:
            book.close()
    if len(table) < 2:
        raise ValueError("Include a header row and at least one player row")
    headers = table[0]
    if any(not h or len(h) > 100 for h in headers) or len(set(h.casefold() for h in headers)) != len(headers):
        raise ValueError("Headers must be non-empty, unique and at most 100 characters")
    rows = []
    for row in table[1:]:
        if len(row) > len(headers):
            raise ValueError("A data row has more cells than the header")
        rows.append(row + [None] * (len(headers) - len(row)))
    suggestions = {}
    used = set()
    for field in IMPORT_FIELDS:
        candidates = [
            h for h in headers if normalized(h) in [field, *ALIASES.get(field, [])] and h not in used
        ]
        if len(candidates) == 1:
            suggestions[field] = candidates[0]
            used.add(candidates[0])
    if formula_count:
        warnings.append(f"{formula_count} formula-like cells withheld; formulas are never evaluated")
    return {
        "headers": headers,
        "rows": rows,
        "source_rows": row_numbers[1:],
        "suggested_mapping": suggestions,
        "warnings": warnings,
        "file_sha256": hashlib.sha256(content).hexdigest(),
        "row_count": len(rows),
        "fields": list(IMPORT_FIELDS),
        "storage": "memory_only",
    }


def number(value: str | None, decimal: str) -> float | None:
    if value is None or value.strip().casefold() in {"", "na", "n/a", "null", "none", "—", "-"}:
        return None
    # Thousands separators are deliberately unsupported: ambiguous input must not change magnitude.
    pattern = r"-?\d+(?:\.\d+)?" if decimal == "." else r"-?\d+(?:,\d+)?"
    if not re.fullmatch(pattern, value.strip()):
        raise ValueError("invalid_number")
    result = float(value.replace(",", "."))
    if not math.isfinite(result) or result < 0 or result > 1_000_000:
        raise ValueError("out_of_range")
    return result


def analyse_import(request: ClubImport, dataset: dict) -> dict:
    canonical = json.dumps(
        request.model_dump(exclude={"label"}), ensure_ascii=False, sort_keys=True, separators=(",", ":")
    )
    import_id = hashlib.sha256(canonical.encode()).hexdigest()[:24]
    index = {field: request.headers.index(column) for field, column in request.mapping.items()}
    source_players = {p["player_id"]: p for p in dataset["players"]}
    reference_profiles = {}
    counts = Counter()
    players = []
    valid_checks = possible_checks = 0
    source_rows = request.source_rows or list(range(2, len(request.rows) + 2))
    for line, row in zip(source_rows, request.rows, strict=True):
        issues = []

        def get(field):
            value = row[index[field]] if field in index else None
            return value.strip() or None if value is not None else None

        name = get("player_name")
        pos = (get("position_group") or "UNK").upper()
        if pos not in POSITIONS:
            pos = "UNK"
            issues.append("position_group:unknown_position")
        if not name or name.startswith(("=", "+", "@")):
            issues.append("player_name:missing_or_invalid")
            name = None
        values = {}
        for field in ("minutes", "age", *DEFINITIONS):
            try:
                value = number(get(field), request.decimal)
                if value is not None and field == "age" and (not value.is_integer() or not 15 <= value <= 60):
                    raise ValueError("out_of_range")
                if (
                    value is not None
                    and field not in {"minutes", "age", "xg", "npxg", "xa"}
                    and not value.is_integer()
                ):
                    raise ValueError("count_must_be_integer")
                if value is not None and field == "minutes" and value > 20_000:
                    raise ValueError("out_of_range")
                values[field] = value
                if value is None and field in index and field != "minutes":
                    issues.append(f"{field}:unavailable")
            except ValueError as exc:
                issues.append(f"{field}:{exc}")
                values[field] = None
        contradicted = set()
        for smaller, upper_bounds in METRIC_BOUNDS.items():
            for terms in upper_bounds:
                value = values[smaller]
                observed = [values[term] for term in terms]
                if (
                    value is not None
                    and all(v is not None for v in observed)
                    and value > sum(observed) + 1e-8
                ):
                    issues.append(f"{smaller}:exceeds_{'+'.join(terms)}")
                    contradicted.add(smaller)
        # Compare the original parsed observations before withholding anything;
        # otherwise the result depends on validation order. Never clamp/infer.
        for field in contradicted:
            values[field] = None
        minutes = values.pop("minutes")
        age = values.pop("age")
        if minutes == 0 and any((v or 0) > 0 for v in values.values()):
            issues.append("minutes:zero_with_events")
            minutes = None
        if minutes is None:
            issues.append("minutes:unavailable")
        metrics = {
            field + "_per90": value * 90 / minutes if value is not None and minutes else None
            for field, value in values.items()
        }
        complete, attempted = values["completed_passes"], values["passes"]
        metrics["pass_completion"] = (
            complete * 100 / attempted if complete is not None and attempted else None
        )
        coverage = 100 * sum(v is not None for v in metrics.values()) / len(metric_catalog())
        # A pasted source name is not proof of a shared definition or competitive context.
        # External benchmarks are available ONLY for exact, identifiable canonical round-trips.
        reference = source_players.get(get("player_id"))
        verified = bool(
            reference
            and name is not None
            and normalized(name)
            in {normalized(reference["player_name"]), normalized(reference["display_name"])}
            and get("dataset_id") == dataset["manifest"]["dataset_id"]
            and minutes is not None
            and math.isclose(minutes, reference["minutes"], rel_tol=1e-10)
            and pos == reference["position_group"]
            and all(
                values[k] is None
                or (
                    reference["totals"].get(k) is not None
                    and math.isclose(values[k], reference["totals"][k], rel_tol=1e-10, abs_tol=1e-10)
                )
                for k in values
            )
        )
        benchmark = None
        if verified and coverage:
            if reference["player_id"] not in reference_profiles:
                reference_profiles[reference["player_id"]] = player_profile(reference, dataset)
            profile = reference_profiles[reference["player_id"]]
            benchmark = {
                "reference_player_id": reference["player_id"],
                "reference_name": reference["display_name"],
                "cohort_size": profile["cohort"]["size"],
                "eligible": profile["cohort"]["eligible"],
                "percentiles": {
                    m["key"]: m["favorable_percentile"]
                    for m in profile["metrics"]
                    if metrics[m["key"]] is not None
                },
            }
        confidence = (
            evidence_confidence(
                [{**m, "value": metrics.get(m["key"])} for m in profile["metrics"]], minutes, True
            )
            if benchmark and benchmark["eligible"]
            else None
        )
        if get("player_id"):
            counts[get("player_id")] += 1
        # Quality counts name/position/minutes + mapped optional cells; completeness is separate below.
        checks = [name is not None, pos != "UNK", minutes is not None]
        checks += [values[f] is not None for f in DEFINITIONS if f in index]
        if "age" in index:
            checks.append(age is not None)
        valid_checks += sum(checks)
        possible_checks += len(checks)
        players.append(
            {
                "player_id": f"upload:{import_id}:row:{line}",
                "source_player_id": get("player_id"),
                "player_name": name,
                "team_name": get("team_name") or request.label,
                "position_group": pos,
                "minutes": minutes,
                "age": age,
                "totals": values,
                "metrics": metrics,
                "data_coverage": coverage,
                "confidence": confidence,
                "issues": issues,
                "source_row": line,
                "benchmark": benchmark,
            }
        )
    for player in players:
        if counts[player["source_player_id"]] > 1:
            player["issues"].append("player_id:duplicate_source_id")
            player["benchmark"] = None
            player["confidence"] = None
        if player["source_player_id"]:
            possible_checks += 1
            valid_checks += counts[player["source_player_id"]] == 1
    return {
        "import_id": import_id,
        "label": request.label,
        "players": players,
        "data_quality": 100 * valid_checks / possible_checks if possible_checks else 0,
        "quality_checks": {"passed": valid_checks, "total": possible_checks},
        "data_coverage": sum(p["data_coverage"] for p in players) / len(players),
        "quality_method": "Valid name, known position, valid minutes and valid mapped numeric cells / expected checks. Not verification of football truth.",
        "rows_with_issues": sum(bool(p["issues"]) for p in players),
        "external_comparison": "Only unchanged canonical rows from the loaded dataset receive an external benchmark. Matching column names do not establish compatible definitions or context.",
        "reference_dataset_id": dataset["manifest"]["dataset_id"],
        "storage": "memory_only",
    }
