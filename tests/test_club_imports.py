import csv
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from apps.api.app.main import app
from apps.api.app.services.catalog import get_catalog
from core.analytics.imports import analyse_import, number, preview_table
from core.schemas.imports import ClubImport


@pytest.mark.parametrize("file_type", ["csv", "xlsx"])
def test_unmapped_scouting_notes_do_not_block_valid_football_observations(file_type):
    # Scouting notes allow 3,000 characters; exports must remain importable.
    note = 'Revisión QA, "sin valoración deportiva".\n' * 70
    assert 256 < len(note) <= 3000
    rows = [["name", "position", "minutes", "shots", "scout_note"], ["Parser fixture", "W", "45", "2", note]]
    if file_type == "csv":
        output = io.StringIO()
        csv.writer(output).writerows(rows)
        content = output.getvalue().encode("utf-8-sig")
    else:
        book = Workbook()
        for row in rows:
            book.active.append(row)
        output = io.BytesIO()
        book.save(output)
        content = output.getvalue()
    preview = preview_table(content, f"notes.{file_type}")
    assert preview["rows"][0][-1] == note.strip()
    assert "scout_note" not in preview["suggested_mapping"]
    request = ClubImport(
        headers=preview["headers"], rows=preview["rows"], mapping=preview["suggested_mapping"]
    )
    player = analyse_import(request, get_catalog()[0])["players"][0]
    assert player["metrics"]["shots_per90"] == 4
    assert player["confidence"] is None
    assert not player["issues"]


def test_import_auxiliary_text_is_bounded_and_cannot_bypass_mapped_field_limits():
    # Unicode characters count once, consistently in Python and the web guard.
    note = "📝" * 4096
    content = ("name,minutes,notes\nParser fixture,90," + note).encode()
    preview = preview_table(content, "notes.csv")
    request = ClubImport(
        headers=preview["headers"], rows=preview["rows"], mapping=preview["suggested_mapping"]
    )
    assert request.rows[0][-1] == note
    with pytest.raises(ValueError, match="Mapped football fields"):
        ClubImport(
            headers=preview["headers"],
            rows=preview["rows"],
            mapping={"player_name": "notes", "minutes": "minutes"},
        )
    with pytest.raises(ValueError, match="4096 characters"):
        preview_table(content + "📝".encode(), "too-long.csv")
    # A direct JSON request still cannot pass an oversized mapped name.
    response = TestClient(app).post(
        "/my-club/analyze",
        json={
            "headers": ["name"],
            "rows": [["x" * 257]],
            "mapping": {"player_name": "name"},
        },
    )
    assert response.status_code == 422
    assert "x" * 257 not in response.text


def test_csv_unicode_mapping_missing_zero_and_inconsistent_counts():
    data = "Jugador;Posición;Minutos;Tiros;Goles;pases;completed_passes\nPrueba A;W;45;2;1;10;8\nPrueba B;ST;0;0;0;;\nPrueba C;CB;90;2;3;5;6\n".encode()
    preview = preview_table(data, "squad.csv")
    assert preview["suggested_mapping"]["position_group"] == "Posición"
    result = analyse_import(
        ClubImport(headers=preview["headers"], rows=preview["rows"], mapping=preview["suggested_mapping"]),
        get_catalog()[0],
    )
    a, b, c = result["players"]
    assert a["metrics"]["shots_per90"] == 4
    assert a["metrics"]["pass_completion"] == 80
    assert b["metrics"]["goals_per90"] is None
    assert b["totals"]["goals"] == 0
    assert b["totals"]["passes"] is None
    assert c["totals"]["goals"] is None
    assert c["metrics"]["pass_completion"] is None
    assert all(p["confidence"] is None and p["benchmark"] is None for p in result["players"])
    assert 0 < result["data_quality"] < 100


def test_xlsx_formulas_not_evaluated_and_duplicate_headers_rejected():
    book = Workbook()
    book.active.append(["player_name", "minutes", "goals"])
    book.active.append(["Invented parser fixture", 90, "=1+1"])
    output = io.BytesIO()
    book.save(output)
    preview = preview_table(output.getvalue(), "squad.xlsx")
    assert preview["rows"] == [["Invented parser fixture", "90", None]]
    assert "formula" in preview["warnings"][0]
    with pytest.raises(ValueError, match="unique"):
        preview_table(b"name,NAME\nA,B", "squad.csv")
    with pytest.raises(ValueError, match="Supported"):
        preview_table(b"dummy", "squad.xls")


@pytest.mark.parametrize("intermediate", ["invalid", "missing"])
def test_import_subtotals_cannot_exceed_observed_ancestors_when_intermediate_totals_fail(intermediate):
    # These are deliberately inconsistent parser fixtures, never football observations.
    fields = [
        "player_name",
        "position_group",
        "minutes",
        "shots",
        "goals",
        "non_penalty_goals",
        "xg",
        "npxg",
        "passes",
        "completed_passes",
        "progressive_passes",
        "key_passes",
        "xa",
    ]
    middle = ["5", "6", "6", "7"] if intermediate == "invalid" else [None] * 4
    request = ClubImport(
        headers=fields,
        mapping={field: field for field in fields},
        rows=[
            [
                "Inconsistent parser fixture",
                "W",
                "90",
                "2",
                middle[0],
                "4",
                middle[1],
                "4",
                "3",
                middle[2],
                "4",
                middle[3],
                "5",
            ]
        ],
    )
    result = analyse_import(request, get_catalog()[0])
    player = result["players"][0]
    for field, upper in (
        ("non_penalty_goals", "shots"),
        ("npxg", "shots"),
        ("progressive_passes", "passes"),
        ("xa", "passes"),
    ):
        assert player["totals"][field] is None
        assert player["metrics"][field + "_per90"] is None
        assert f"{field}:exceeds_{upper}" in player["issues"]
    assert player["metrics"]["shots_per90"] == 2
    assert player["metrics"]["passes_per90"] == 3
    assert player["data_coverage"] == pytest.approx(200 / 21)
    assert player["benchmark"] is None and player["confidence"] is None


def test_import_unknown_intermediate_does_not_erase_a_compatible_observed_subtotal():
    fields = [
        "player_name",
        "position_group",
        "minutes",
        "shots",
        "non_penalty_goals",
        "npxg",
        "passes",
        "progressive_passes",
        "xa",
    ]
    request = ClubImport(
        headers=fields,
        mapping={field: field for field in fields},
        rows=[["Incomplete parser fixture", "W", "45", "2", "1", "1.5", "10", "3", "0.5"]],
    )
    player = analyse_import(request, get_catalog()[0])["players"][0]
    assert player["metrics"]["non_penalty_goals_per90"] == 2
    assert player["metrics"]["npxg_per90"] == 3
    assert player["metrics"]["progressive_passes_per90"] == 6
    assert player["metrics"]["xa_per90"] == 1
    assert all(player["totals"][key] is None for key in ("goals", "xg", "completed_passes", "key_passes"))
    assert player["metrics"]["pass_completion"] is None
    assert not player["issues"]


def test_import_assists_final_third_and_box_entries_respect_their_event_denominators():
    fields = [
        "player_name",
        "position_group",
        "minutes",
        "passes",
        "completed_passes",
        "carries",
        "assists",
        "passes_into_final_third",
        "entries_into_box",
    ]
    request = ClubImport(
        headers=fields,
        mapping={field: field for field in fields},
        rows=[
            ["Inconsistent parser fixture", "W", "90", "3", None, "2", "4", "4", "6"],
            ["Incomplete parser fixture", "W", "90", None, None, "2", None, None, "6"],
            ["Compatible parser fixture", "W", "45", "3", "2", "2", "1", "2", "4"],
        ],
    )
    bad, unknown, good = analyse_import(request, get_catalog()[0])["players"]
    assert all(
        bad["metrics"][field + "_per90"] is None
        for field in ("assists", "passes_into_final_third", "entries_into_box")
    )
    assert "entries_into_box:exceeds_passes+carries" in bad["issues"]
    # Unknown pass counts cannot be treated as zero in a pass+carry upper bound.
    assert unknown["metrics"]["entries_into_box_per90"] == 6
    assert good["metrics"]["entries_into_box_per90"] == 8
    assert good["metrics"]["passes_into_final_third_per90"] == 4
    assert good["metrics"]["assists_per90"] == 2


def test_malformed_lazy_xlsx_cell_is_a_controlled_upload_error():
    book = Workbook()
    book.active.append(["name", "minutes"])
    book.active.append(["Parser fixture", 90])
    original = io.BytesIO()
    book.save(original)
    damaged = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(damaged, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(
                    b'<c r="B2" t="n"><v>90</v></c>',
                    b'<c r="B2" t="s"><v>99</v></c>',
                )
            target.writestr(item.filename, data)
    with pytest.raises(ValueError, match="Invalid XLSX cell"):
        preview_table(damaged.getvalue(), "fixture.xlsx")
    response = TestClient(app).post("/my-club/preview?filename=fixture.xlsx", content=damaged.getvalue())
    assert response.status_code == 422
    assert "Parser fixture" not in response.text


def test_import_bounds_decimal_and_duplicate_ids():
    assert number("90,5", ",") == 90.5
    for text in ("1,000", "NaN", "inf", "1e4", "-1", "=2+2"):
        with pytest.raises(ValueError):
            number(text, ".")
    with pytest.raises(ValueError, match="60 columns"):
        preview_table((",".join(f"col{i}" for i in range(61)) + "\n1,2").encode(), "s.csv")
    payload = ClubImport(
        headers=["name", "id", "position", "minutes"],
        rows=[["A", "same", "W", "90"], ["B", "same", "W", "90"]],
        mapping={
            "player_name": "name",
            "player_id": "id",
            "position_group": "position",
            "minutes": "minutes",
        },
    )
    result = analyse_import(payload, get_catalog()[0])
    assert result["players"][0]["player_id"] != result["players"][1]["player_id"]
    assert result["import_id"] == analyse_import(payload, get_catalog()[0])["import_id"]
    assert all("player_id:duplicate_source_id" in p["issues"] for p in result["players"])
    assert result["data_quality"] < 100


def test_only_unchanged_canonical_rows_get_external_benchmarks():
    dataset = get_catalog()[0]
    source = next(p for p in dataset["players"] if p["display_name"] == "Lamine Yamal")
    fields = ["player_name", "player_id", "dataset_id", "position_group", "minutes", "shots"]
    values = [
        source["display_name"],
        source["player_id"],
        source["dataset_id"],
        source["position_group"],
        str(source["minutes"]),
        str(source["totals"]["shots"]),
    ]
    request = ClubImport(headers=fields, rows=[values], mapping=dict(zip(fields, fields)))
    p = analyse_import(request, dataset)["players"][0]
    assert p["benchmark"]["reference_player_id"] == source["player_id"]
    assert 0 < p["confidence"] < 10  # one observed metric, not the source's full coverage
    request.rows[0][-1] = "99"
    assert analyse_import(request, dataset)["players"][0]["benchmark"] is None


def test_quality_exposes_exact_counts_separately_from_coverage_and_confidence():
    fields = ["player_name", "player_id", "position_group", "minutes", "shots", "goals", "age"]
    request = ClubImport(
        headers=fields,
        mapping={field: field for field in fields},
        rows=[
            ["Parser A", "duplicate", "W", "90", "2", "1", "23"],
            ["Parser B", "duplicate", "unknown", "bad", "3", "7", "bad"],
        ],
    )
    result = analyse_import(request, get_catalog()[0])
    # Each row: name/position/minutes + mapped shots/goals/age + source-ID uniqueness.
    # First: six valid, duplicate ID fails. Second: only name and shots are valid.
    assert result["quality_checks"] == {"passed": 8, "total": 14}
    assert result["data_quality"] == pytest.approx(100 * 8 / 14)
    assert result["players"][0]["data_coverage"] == pytest.approx(200 / 21)
    assert result["players"][1]["data_coverage"] == 0
    assert all(p["confidence"] is None for p in result["players"])


def test_real_http_upload_analysis_and_validation_errors():
    client = TestClient(app)
    preview = client.post(
        "/my-club/preview?filename=squad.csv", content=b"name,minutes,position,goals\nInvented test,90,ST,2"
    ).json()
    response = client.post(
        "/my-club/analyze",
        json={
            "headers": preview["headers"],
            "rows": preview["rows"],
            "mapping": preview["suggested_mapping"],
        },
    )
    assert response.status_code == 200
    assert response.json()["players"][0]["metrics"]["goals_per90"] == 2
    assert client.post("/my-club/preview", content=b"x" * (3 * 1024 * 1024 + 1)).status_code == 413
    bad = client.post(
        "/my-club/analyze",
        json={"headers": ["name"], "rows": [["private name"]], "mapping": {"unknown": "name"}},
    )
    assert bad.status_code == 422
    assert "private name" not in bad.text


def test_original_file_row_lineage_survives_blank_rows_in_csv_and_xlsx():
    csv_data = b"\nname,position,minutes,shots\nFirst fixture,W,45,2\n\nSecond fixture,ST,90,3\n"
    book = Workbook()
    for row in [
        [],
        ["name", "position", "minutes", "shots"],
        ["First fixture", "W", 45, 2],
        [],
        ["Second fixture", "ST", 90, 3],
    ]:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    for content, name in [(csv_data, "fixture.csv"), (output.getvalue(), "fixture.xlsx")]:
        preview = preview_table(content, name)
        assert preview["source_rows"] == [3, 5]
        request = ClubImport(
            headers=preview["headers"],
            rows=preview["rows"],
            source_rows=preview["source_rows"],
            mapping=preview["suggested_mapping"],
        )
        players = analyse_import(request, get_catalog()[0])["players"]
        assert [p["source_row"] for p in players] == [3, 5]
        assert players[1]["player_id"].endswith(":row:5")
        assert players[0]["metrics"]["shots_per90"] == 4
    with pytest.raises(ValueError, match="Source row numbers"):
        ClubImport(headers=["name"], rows=[["A"], ["B"]], source_rows=[2, 2], mapping={"player_name": "name"})
